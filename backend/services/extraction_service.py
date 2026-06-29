from pathlib import Path

import fitz
from pptx import Presentation

from models.uploaded_file_model import PageRef


def extract_text_from_file(storage_path: str, file_type: str) -> tuple[str, list[PageRef]]:
    if file_type == "pdf":
        page_refs = extract_pdf_text(storage_path)
    elif file_type == "pptx":
        page_refs = extract_pptx_text(storage_path)
    else:
        raise ValueError("Unsupported file type")

    extracted_text = "\n\n".join(page["text"] for page in page_refs)
    return extracted_text, page_refs


def extract_pdf_text(storage_path: str) -> list[PageRef]:
    page_refs: list[PageRef] = []

    with fitz.open(storage_path) as document:
        for page_index, page in enumerate(document, start=1):
            page_refs.append(
                {
                    "page": page_index,
                    "text": page.get_text().strip(),
                }
            )

    return page_refs


def extract_pptx_text(storage_path: str) -> list[PageRef]:
    presentation = Presentation(storage_path)
    page_refs: list[PageRef] = []

    for slide_index, slide in enumerate(presentation.slides, start=1):
        text_parts: list[str] = []

        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text:
                text_parts.append(shape.text.strip())

        page_refs.append(
            {
                "page": slide_index,
                "text": "\n".join(text_parts).strip(),
            }
        )

    return page_refs


def build_storage_path(
    user_id: str,
    course_id: str,
    upload_id: str,
    original_filename: str,
) -> Path:
    safe_filename = Path(original_filename).name.replace(" ", "_")
    return Path("uploads") / user_id / course_id / f"{upload_id}_{safe_filename}"
